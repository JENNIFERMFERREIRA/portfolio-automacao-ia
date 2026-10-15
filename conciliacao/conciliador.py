#!/usr/bin/env python3
"""Conciliação bancária automatizada.

Cruza os boletos recebidos (exportados do banco em CSV) com a planilha de controle de pagamentos
e gera um relatório em Excel dizendo o que confere, o que tem valor diferente e o que falta lançar.

Destaques técnicos:
  * comparação "fuzzy" de nomes: o banco abrevia o pagador ("MARIA S OLIVEIRA") e a planilha tem o
    nome completo ("Maria da Silva Oliveira") — tokens normalizados, sem acento e stopwords;
  * divisão dos boletos entre os créditos que caíram no extrato (problema da soma de subconjuntos,
    resolvido com programação dinâmica em centavos);
  * saída em Excel formatada com openpyxl.

Uso:
    python conciliador.py --boletos exemplo/boletos.csv --controle exemplo/controle.xlsx \
                          --creditos 1520.40,830.10 --saida relatorio.xlsx

Todos os dados da pasta exemplo/ são fictícios.
"""
import argparse
import csv
import difflib
import itertools
import re
import unicodedata

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

STOPWORDS = {"DA", "DE", "DO", "DAS", "DOS", "E"}


def tokens(nome: str) -> list[str]:
    """'Maria da Silva-Oliveira' -> ['MARIA', 'SILVA', 'OLIVEIRA']"""
    s = unicodedata.normalize("NFKD", str(nome or "")).encode("ascii", "ignore").decode().upper()
    return [t for t in re.sub(r"[^A-Z ]", " ", s).split() if t not in STOPWORDS]


def token_igual(abrev: str, completo: str) -> bool:
    """Aceita iniciais ('S' == 'SILVA'), prefixos e pequenos erros de digitação."""
    if abrev == completo:
        return True
    if len(abrev) <= 2:
        return completo.startswith(abrev)
    return completo.startswith(abrev) or difflib.SequenceMatcher(None, abrev, completo).ratio() >= 0.8


def pontuacao(abreviado: str, completo: str) -> int:
    """0 = não é a mesma pessoa. Quanto maior, mais tokens bateram (na ordem)."""
    a, c = tokens(abreviado), tokens(completo)
    if not a or not c or not token_igual(a[0], c[0]):
        return 0
    pos, acertos = 0, 0
    for t in a[1:]:
        while pos < len(c) - 1:
            pos += 1
            if token_igual(t, c[pos]):
                acertos += 1
                break
        else:
            return 0
    return 1 + acertos


def dividir_por_credito(boletos: list[dict], creditos: list[float]) -> list[list[dict]]:
    """Separa os boletos em grupos cuja soma líquida bate exatamente com cada crédito do extrato.
    Soma de subconjuntos por programação dinâmica (valores em centavos para evitar erro de float)."""
    cent = lambda b: round((b["valor"] - b["tarifa"]) * 100)
    alvos = [round(c * 100) for c in creditos]
    if sum(alvos) != sum(cent(b) for b in boletos):
        raise SystemExit(f"Os créditos {creditos} não somam o líquido dos boletos.")
    if len(alvos) == 1:
        return [boletos]
    for ordem in itertools.permutations(range(len(alvos))):
        restantes, grupos = list(boletos), {}
        for k in ordem[:-1]:
            alcancavel = {0: None}            # soma -> (índice do boleto, soma anterior)
            for i, b in enumerate(restantes):
                for s in list(alcancavel):
                    ns = s + cent(b)
                    if ns <= alvos[k] and ns not in alcancavel:
                        alcancavel[ns] = (i, s)
            if alvos[k] not in alcancavel:
                break
            usados, s = set(), alvos[k]
            while s:
                i, s = alcancavel[s]
                usados.add(i)
            grupos[k] = [b for i, b in enumerate(restantes) if i in usados]
            restantes = [b for i, b in enumerate(restantes) if i not in usados]
        else:
            grupos[ordem[-1]] = restantes
            return [grupos[k] for k in range(len(alvos))]
    raise SystemExit("Nenhuma combinação de boletos fecha os créditos informados.")


def ler_controle(caminho: str) -> list[dict]:
    """Planilha de controle: colunas VALOR | DATA | CLIENTE | CONFERIDO"""
    ws = openpyxl.load_workbook(caminho, data_only=True).active
    linhas = []
    for n, (valor, data, cliente, conferido) in enumerate(ws.iter_rows(min_row=2, max_col=4, values_only=True), 2):
        if cliente and isinstance(valor, (int, float)):
            linhas.append(dict(linha=n, valor=round(valor, 2), data=data, cliente=str(cliente).strip(),
                               conferido=bool(conferido)))
    return linhas


def conciliar(boletos: list[dict], controle: list[dict]) -> list[dict]:
    usados = set()
    for b in boletos:
        candidatos = [(pontuacao(b["pagador"], c["cliente"]), c) for c in controle if c["linha"] not in usados]
        candidatos = [(p, c) for p, c in candidatos if p]
        mesmo_valor = [x for x in candidatos if abs(x[1]["valor"] - b["valor"]) < 0.01]
        escolha = sorted(mesmo_valor or candidatos, key=lambda x: -x[0])
        if not escolha:
            b.update(situacao="SEM LANÇAMENTO NA PLANILHA", cliente=None, linha=None)
            continue
        c = escolha[0][1]
        usados.add(c["linha"])
        if abs(c["valor"] - b["valor"]) >= 0.01:
            situacao = "VALOR DIFERENTE"
        elif c["conferido"]:
            situacao = "JÁ CONFERIDO (possível duplicidade)"
        else:
            situacao = "CONFERE"
        b.update(situacao=situacao, cliente=c["cliente"], linha=c["linha"], valor_planilha=c["valor"])
    return boletos


def salvar(boletos, grupos, creditos, caminho):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Conciliação"
    cab = ["PAGADOR (BANCO)", "CLIENTE (PLANILHA)", "LINHA", "VALOR", "TARIFA", "LÍQUIDO", "CRÉDITO", "SITUAÇÃO"]
    ws.append(cab)
    for cel in ws[1]:
        cel.font = Font(bold=True, color="FFFFFF")
        cel.fill = PatternFill("solid", fgColor="1F4E78")
    cores = {"CONFERE": "E2EFDA", "VALOR DIFERENTE": "FFF2CC"}
    for credito, grupo in zip(creditos, grupos):
        for b in grupo:
            ws.append([b["pagador"], b["cliente"] or "—", b["linha"] or "—", b["valor"], b["tarifa"],
                       round(b["valor"] - b["tarifa"], 2), credito, b["situacao"]])
            cor = cores.get(b["situacao"], "F8CBAD")
            for cel in ws[ws.max_row]:
                cel.fill = PatternFill("solid", fgColor=cor)
                cel.alignment = Alignment(vertical="top")
    for col, largura in zip("ABCDEFGH", [26, 30, 7, 10, 8, 10, 10, 36]):
        ws.column_dimensions[col].width = largura
    wb.save(caminho)


def main():
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--boletos", required=True, help="CSV: pagador,valor,tarifa,data")
    ap.add_argument("--controle", required=True, help="planilha de controle (.xlsx)")
    ap.add_argument("--creditos", default="", help="créditos do extrato, ex.: 1520.40,830.10")
    ap.add_argument("--saida", default="relatorio_conciliacao.xlsx")
    a = ap.parse_args()

    with open(a.boletos, encoding="utf-8") as f:
        boletos = [dict(pagador=r["pagador"].strip().upper(), valor=float(r["valor"]),
                        tarifa=float(r["tarifa"]), data=r["data"]) for r in csv.DictReader(f)]
    boletos = conciliar(boletos, ler_controle(a.controle))
    liquido = round(sum(b["valor"] - b["tarifa"] for b in boletos), 2)
    creditos = [float(x) for x in a.creditos.split(",") if x.strip()] or [liquido]
    grupos = dividir_por_credito(boletos, creditos)
    salvar(boletos, grupos, creditos, a.saida)

    for b in boletos:
        print(f"{b['pagador']:<24} R$ {b['valor']:>8.2f}  {b['situacao']:<36} {b['cliente'] or ''}")
    print(f"\nLíquido: R$ {liquido:.2f} em {len(creditos)} crédito(s) -> {a.saida}")


if __name__ == "__main__":
    main()
